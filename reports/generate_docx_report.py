import os
from pathlib import Path

def main():
    print("[*] Generating formatted Project Report...")
    reports_dir = Path(__file__).resolve().parent
    report_md = reports_dir / "PROJECT_REPORT.md"
    output_docx = reports_dir / "PROJECT_REPORT.docx"
    output_pdf = reports_dir / "PROJECT_REPORT.pdf"

    if not report_md.exists():
        print(f"[-] Error: Source report {report_md} not found!")
        return

    try:
        # Check if python-docx is installed
        import docx
        doc = docx.Document()
        doc.add_heading("Stitch Sentinel - Windows EDR Project Report", 0)
        
        with open(report_md, "r", encoding="utf-8") as f:
            lines = f.readlines()
            
        for line in lines:
            line_str = line.strip()
            if line_str.startswith("# "):
                doc.add_heading(line_str[2:], level=1)
            elif line_str.startswith("## "):
                doc.add_heading(line_str[3:], level=2)
            elif line_str.startswith("### "):
                doc.add_heading(line_str[4:], level=3)
            elif line_str:
                doc.add_paragraph(line_str)
                
        doc.save(output_docx)
        print(f"[+] Successfully generated DOCX report at: {output_docx}")
    except ImportError:
        # Fallback to basic text report copy
        print("[!] python-docx library not installed. Creating plain-text replica...")
        output_txt = reports_dir / "PROJECT_REPORT.txt"
        with open(report_md, "r", encoding="utf-8") as src, open(output_txt, "w", encoding="utf-8") as dst:
            dst.write(src.read())
        print(f"[+] Successfully generated plain-text report replica at: {output_txt}")

if __name__ == "__main__":
    main()
